from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Literal
from uuid import UUID, uuid4

from docx import Document
from pypdf import PdfWriter
import pytest
from sqlalchemy import create_engine, event, func, insert, select
from sqlalchemy.engine import Engine

from app.config import Settings
from app.domain.adventure_imports.errors import (
    AdventureImportForbiddenError,
    AdventureImportNotFoundError,
    AdventureImportValidationError,
    ExtractorUnavailableError,
)
from app.domain.adventure_imports.extractors import (
    ExtractionResult,
    extract_docx_text,
    extract_pdf_text,
)
from app.domain.adventure_imports.service import AdventureImportService
from app.domain.room_assets.schemas import RoomAsset
from app.domain.room_assets.service import RoomAssetService
from app.domain.rooms.schemas import RoomAccessAuthority, RoomAccessContext
from app.persistence.adventure_imports.repository import AdventureImportRepository
from app.persistence.adventure_imports.tables import (
    adventure_import_drafts,
    adventure_import_sources,
    adventure_imports,
)
from app.persistence.adventures.tables import adventure_definitions
from app.persistence.room_assets.repository import RoomAssetRepository
from app.persistence.room_assets.storage import FilesystemAssetStorage
from app.persistence.room_assets.tables import room_assets
from app.persistence.rooms.tables import campaigns, rooms

TABLES_TO_CREATE = [
    rooms,
    campaigns,
    adventure_definitions,
    room_assets,
    adventure_imports,
    adventure_import_sources,
    adventure_import_drafts,
]


def make_pdf(pages_text: list[str]) -> bytes:
    """Generate a minimal valid PDF with the given text on each page."""
    objs: list[bytes] = []
    # 1: Catalog
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    # 2: Pages
    kids = " ".join(f"{3 + i} 0 R" for i in range(len(pages_text)))
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages_text)} >>".encode("ascii"))

    font_obj_idx = 3 + len(pages_text) * 2
    for i in range(len(pages_text)):
        contents_idx = 3 + len(pages_text) + i
        page_dict = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {contents_idx} 0 R /Resources << /Font << /F1 {font_obj_idx} 0 R >> >> >>"
        )
        objs.append(page_dict.encode("ascii"))

    for text in pages_text:
        stream_data = f"BT /F1 12 Tf 100 700 Td ({text}) Tj ET".encode("ascii")
        content_dict = (
            f"<< /Length {len(stream_data)} >>\nstream\n".encode("ascii")
            + stream_data
            + b"\nendstream"
        )
        objs.append(content_dict)

    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = [b"%PDF-1.4\n"]
    offsets = [0]
    for i, obj in enumerate(objs, 1):
        offsets.append(sum(len(x) for x in out))
        out.append(f"{i} 0 obj\n".encode("ascii") + obj + b"\nendobj\n")

    xref_offset = sum(len(x) for x in out)
    out.append(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode("ascii"))
    for off in offsets[1:]:
        out.append(f"{off:010d} 00000 n \n".encode("ascii"))
    out.append(
        f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode(
            "ascii"
        )
    )
    return b"".join(out)


def make_empty_pdf() -> bytes:
    """Generate a PDF with a blank page containing no text."""
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = BytesIO()
    writer.write(buf)
    return buf.getvalue()


def make_docx(paragraphs: list[tuple[str, str]]) -> bytes:
    """Generate a DOCX document with (text, style_name) paragraphs."""
    doc = Document()
    for text, style_name in paragraphs:
        if style_name.lower().startswith("heading"):
            level = 1
            if style_name.split()[-1].isdigit():
                level = int(style_name.split()[-1])
            doc.add_heading(text, level=level)
        else:
            doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_empty_docx() -> bytes:
    """Generate an empty DOCX document with no paragraphs."""
    doc = Document()
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


@dataclass(frozen=True)
class ExtractionFixture:
    engine: Engine
    import_repo: AdventureImportRepository
    asset_repo: RoomAssetRepository
    asset_storage: FilesystemAssetStorage
    asset_service: RoomAssetService
    import_service: AdventureImportService
    settings: Settings
    room_a_id: UUID
    room_b_id: UUID
    owner_a: RoomAccessContext
    dm_a: RoomAccessContext
    member_a: RoomAccessContext
    owner_b: RoomAccessContext


def _create_sqlite_engine(db_path: Path) -> Engine:
    engine = create_engine(f"sqlite+pysqlite:///{db_path.as_posix()}")

    @event.listens_for(engine, "connect")
    def _enable_fks(dbapi_conn, _connection_record) -> None:
        cursor = dbapi_conn.cursor()
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()

    for table in TABLES_TO_CREATE:
        table.create(engine, checkfirst=True)
    return engine


@pytest.fixture
def fix(tmp_path: Path) -> ExtractionFixture:
    db_path = tmp_path / "test_extraction.db"
    engine = _create_sqlite_engine(db_path)
    now = datetime.now(timezone.utc)
    room_a_id = uuid4()
    room_b_id = uuid4()

    with engine.begin() as conn:
        conn.execute(
            insert(rooms).values(
                [
                    {
                        "id": room_a_id,
                        "code": "ROOM-A",
                        "name": "Room A",
                        "password_salt": b"salt",
                        "password_hash": b"pw",
                        "owner_key_hash": b"owner",
                        "dm_key_hash": b"dm",
                        "created_at": now,
                        "updated_at": now,
                    },
                    {
                        "id": room_b_id,
                        "code": "ROOM-B",
                        "name": "Room B",
                        "password_salt": b"salt",
                        "password_hash": b"pw",
                        "owner_key_hash": b"owner",
                        "dm_key_hash": b"dm",
                        "created_at": now,
                        "updated_at": now,
                    },
                ]
            )
        )

    settings = Settings()
    import_repo = AdventureImportRepository(engine)
    asset_storage = FilesystemAssetStorage(tmp_path / "assets")
    asset_repo = RoomAssetRepository(engine)
    asset_service = RoomAssetService(
        asset_repo,
        asset_storage,
        max_image_bytes=settings.asset_max_image_bytes,
        max_source_document_bytes=settings.asset_max_source_document_bytes,
    )
    import_service = AdventureImportService(import_repo, settings, asset_service)

    return ExtractionFixture(
        engine=engine,
        import_repo=import_repo,
        asset_repo=asset_repo,
        asset_storage=asset_storage,
        asset_service=asset_service,
        import_service=import_service,
        settings=settings,
        room_a_id=room_a_id,
        room_b_id=room_b_id,
        owner_a=RoomAccessContext(
            room_id=room_a_id,
            access_session_id=uuid4(),
            authority=RoomAccessAuthority.OWNER,
            display_name="Owner A",
        ),
        dm_a=RoomAccessContext(
            room_id=room_a_id,
            access_session_id=uuid4(),
            authority=RoomAccessAuthority.DM,
            display_name="DM A",
        ),
        member_a=RoomAccessContext(
            room_id=room_a_id,
            access_session_id=uuid4(),
            authority=RoomAccessAuthority.MEMBER,
            display_name="Member A",
        ),
        owner_b=RoomAccessContext(
            room_id=room_b_id,
            access_session_id=uuid4(),
            authority=RoomAccessAuthority.OWNER,
            display_name="Owner B",
        ),
    )


def _count_import_rows(engine: Engine) -> tuple[int, int, int]:
    with engine.connect() as conn:
        c1 = conn.execute(select(func.count()).select_from(adventure_imports)).scalar_one()
        c2 = conn.execute(select(func.count()).select_from(adventure_import_sources)).scalar_one()
        c3 = conn.execute(select(func.count()).select_from(adventure_import_drafts)).scalar_one()
        return (c1, c2, c3)


# --- 1. Deterministic Extraction and Exact Locators ---


def test_pdf_extraction_deterministic_locators(fix: ExtractionFixture) -> None:
    page_0_text = "Chapter 1: The Dark Forest"
    page_1_text = "Chapter 2: The Hidden Cave"
    pdf_bytes = make_pdf([page_0_text, page_1_text])

    # Direct extractor unit test
    result = extract_pdf_text(pdf_bytes)
    assert len(result.sections) == 2
    s0 = result.sections[0]
    s1 = result.sections[1]
    assert s0["page_index"] == 0
    assert s0["start_offset"] == 0
    assert s0["end_offset"] == len(page_0_text)
    assert result.normalized_text[int(s0["start_offset"]) : int(s0["end_offset"])] == page_0_text

    assert s1["page_index"] == 1
    assert s1["start_offset"] == len(page_0_text) + 2
    assert s1["end_offset"] == len(page_0_text) + 2 + len(page_1_text)
    assert result.normalized_text[int(s1["start_offset"]) : int(s1["end_offset"])] == page_1_text

    # Upload as room asset
    asset = fix.asset_service.create(
        fix.owner_a,
        room_id=fix.room_a_id,
        kind="source_document",
        filename="adventure.pdf",
        mime_type="application/pdf",
        data=pdf_bytes,
        visibility="dm_only",
    )

    imp = fix.import_service.create_import(fix.owner_a, fix.room_a_id, "PDF Import")
    src = fix.import_service.add_asset_source(
        fix.owner_a, fix.room_a_id, imp.id, asset_id=asset.id
    )

    assert src.source_kind == "pdf"
    assert src.asset_id == asset.id
    assert src.metadata_json["filename"] == "adventure.pdf"
    assert src.metadata_json["media_type"] == "application/pdf"
    assert src.metadata_json["byte_size"] == len(pdf_bytes)
    assert src.metadata_json["line_count"] == 3
    assert src.metadata_json["char_count"] == len(result.normalized_text)
    assert src.metadata_json["sections"] == list(result.sections)

    # Chunk reading matches locator offsets exactly
    chunk = fix.import_service.read_source_chunk(
        fix.owner_a, fix.room_a_id, imp.id, src.id
    )
    assert chunk.text[int(s0["start_offset"]) : int(s0["end_offset"])] == page_0_text
    assert chunk.text[int(s1["start_offset"]) : int(s1["end_offset"])] == page_1_text

    # Clean draft with no warnings
    draft = fix.import_service.get_draft(fix.owner_a, fix.room_a_id, imp.id)
    assert draft.warnings == []


def test_docx_extraction_deterministic_locators(fix: ExtractionFixture) -> None:
    doc_paragraphs = [
        ("The Lost Mine", "Heading 1"),
        ("This is the introduction paragraph.", "Normal"),
        ("Goblin Ambush", "Heading 2"),
        ("Arrows fly from the thicket.", "Normal"),
    ]
    docx_bytes = make_docx(doc_paragraphs)

    # Direct extractor unit test
    result = extract_docx_text(docx_bytes)
    assert len(result.sections) == 4

    # Heading monotonic index check
    assert result.sections[0]["paragraph_index"] == 0
    assert result.sections[0]["heading_index"] == 0
    assert result.sections[1]["paragraph_index"] == 1
    assert result.sections[1]["heading_index"] is None
    assert result.sections[2]["paragraph_index"] == 2
    assert result.sections[2]["heading_index"] == 1
    assert result.sections[3]["paragraph_index"] == 3
    assert result.sections[3]["heading_index"] is None

    # Exact locator slicing into normalized text
    for i, (text, _) in enumerate(doc_paragraphs):
        sec = result.sections[i]
        start = int(sec["start_offset"])
        end = int(sec["end_offset"])
        assert result.normalized_text[start:end] == text

    # Upload as room asset
    asset = fix.asset_service.create(
        fix.owner_a,
        room_id=fix.room_a_id,
        kind="source_document",
        filename="module.docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        data=docx_bytes,
        visibility="dm_only",
    )

    imp = fix.import_service.create_import(fix.owner_a, fix.room_a_id, "DOCX Import")
    src = fix.import_service.add_asset_source(
        fix.owner_a, fix.room_a_id, imp.id, asset_id=asset.id
    )

    assert src.source_kind == "docx"
    assert src.asset_id == asset.id
    assert src.metadata_json["filename"] == "module.docx"
    assert src.metadata_json["sections"] == list(result.sections)

    # Clean draft with no warnings
    draft = fix.import_service.get_draft(fix.owner_a, fix.room_a_id, imp.id)
    assert draft.warnings == []


# --- 2. Empty Extraction Warning and Empty State ---


@pytest.mark.parametrize(
    "source_kind,mime_type,data_factory",
    [
        (
            "pdf",
            "application/pdf",
            make_empty_pdf,
        ),
        (
            "docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            make_empty_docx,
        ),
    ],
)
def test_empty_extraction_creates_warning_and_empty_state(
    fix: ExtractionFixture,
    source_kind: str,
    mime_type: str,
    data_factory: callable,
) -> None:
    data = data_factory()
    asset = fix.asset_service.create(
        fix.owner_a,
        room_id=fix.room_a_id,
        kind="source_document",
        filename=f"empty.{source_kind}",
        mime_type=mime_type,
        data=data,
        visibility="dm_only",
    )

    imp = fix.import_service.create_import(fix.owner_a, fix.room_a_id, f"Empty {source_kind.upper()}")
    src = fix.import_service.add_asset_source(
        fix.owner_a, fix.room_a_id, imp.id, asset_id=asset.id
    )

    assert src.source_kind == source_kind
    assert src.text_length == 0
    assert src.metadata_json["char_count"] == 0
    assert src.metadata_json["sections"] == []

    # Import status remains "source"
    imp_stored = fix.import_service.get_import(fix.owner_a, fix.room_a_id, imp.id)
    assert imp_stored.status == "source"

    # Exactly one stable warning appended to draft
    draft = fix.import_service.get_draft(fix.owner_a, fix.room_a_id, imp.id)
    assert len(draft.warnings) == 1
    w = draft.warnings[0]
    assert w.warning_id == f"source:{src.id}:empty_extraction"
    assert w.level == "warning"
    assert w.code == "empty_extraction"
    assert w.source_id == src.id


# --- 3. Idempotency ---


def test_repeated_same_asset_is_idempotent_with_one_warning(fix: ExtractionFixture) -> None:
    empty_pdf_bytes = make_empty_pdf()
    asset = fix.asset_service.create(
        fix.owner_a,
        room_id=fix.room_a_id,
        kind="source_document",
        filename="empty.pdf",
        mime_type="application/pdf",
        data=empty_pdf_bytes,
        visibility="dm_only",
    )

    imp = fix.import_service.create_import(fix.owner_a, fix.room_a_id, "Idempotent Import")
    src_1 = fix.import_service.add_asset_source(
        fix.owner_a, fix.room_a_id, imp.id, asset_id=asset.id
    )
    counts_after_1 = _count_import_rows(fix.engine)

    # Re-adding the exact same asset
    src_2 = fix.import_service.add_asset_source(
        fix.owner_a, fix.room_a_id, imp.id, asset_id=asset.id
    )
    counts_after_2 = _count_import_rows(fix.engine)

    assert src_1.id == src_2.id
    assert counts_after_1 == counts_after_2

    draft = fix.import_service.get_draft(fix.owner_a, fix.room_a_id, imp.id)
    assert len(draft.warnings) == 1


def test_distinct_assets_with_identical_bytes_remain_distinct_sources(
    fix: ExtractionFixture,
) -> None:
    pdf_bytes = make_pdf(["Shared content across files"])
    asset_1 = fix.asset_service.create(
        fix.owner_a,
        room_id=fix.room_a_id,
        kind="source_document",
        filename="doc1.pdf",
        mime_type="application/pdf",
        data=pdf_bytes,
        visibility="dm_only",
    )
    asset_2 = fix.asset_service.create(
        fix.owner_a,
        room_id=fix.room_a_id,
        kind="source_document",
        filename="doc2.pdf",
        mime_type="application/pdf",
        data=pdf_bytes,
        visibility="dm_only",
    )
    assert asset_1.id != asset_2.id

    imp = fix.import_service.create_import(fix.owner_a, fix.room_a_id, "Distinct Assets")
    src_1 = fix.import_service.add_asset_source(
        fix.owner_a, fix.room_a_id, imp.id, asset_id=asset_1.id
    )
    src_2 = fix.import_service.add_asset_source(
        fix.owner_a, fix.room_a_id, imp.id, asset_id=asset_2.id
    )

    assert src_1.id != src_2.id
    sources = fix.import_service.list_sources(fix.owner_a, fix.room_a_id, imp.id)
    assert len(sources) == 2


# --- 4. Unavailable Library Path ---


def test_unavailable_library_creates_warning_not_500(
    fix: ExtractionFixture,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pdf_bytes = make_pdf(["Some text"])
    asset = fix.asset_service.create(
        fix.owner_a,
        room_id=fix.room_a_id,
        kind="source_document",
        filename="doc.pdf",
        mime_type="application/pdf",
        data=pdf_bytes,
        visibility="dm_only",
    )

    def _mock_unavailable(_data: bytes) -> ExtractionResult:
        raise ExtractorUnavailableError("pypdf is not installed")

    monkeypatch.setattr("app.domain.adventure_imports.service.extract_pdf_text", _mock_unavailable)

    imp = fix.import_service.create_import(fix.owner_a, fix.room_a_id, "Unavailable Import")
    src = fix.import_service.add_asset_source(
        fix.owner_a, fix.room_a_id, imp.id, asset_id=asset.id
    )

    assert src.text_length == 0
    assert src.metadata_json["sections"] == []

    draft = fix.import_service.get_draft(fix.owner_a, fix.room_a_id, imp.id)
    assert len(draft.warnings) == 1
    w = draft.warnings[0]
    assert w.warning_id == f"source:{src.id}:extractor_unavailable"
    assert w.level == "warning"
    assert w.code == "extractor_unavailable"
    assert w.source_id == src.id

    # Re-adding stays idempotent
    counts_before = _count_import_rows(fix.engine)
    src_again = fix.import_service.add_asset_source(
        fix.owner_a, fix.room_a_id, imp.id, asset_id=asset.id
    )
    counts_after = _count_import_rows(fix.engine)
    assert src_again.id == src.id
    assert counts_before == counts_after


# --- 5. Rejections and Zero Side Effects ---


def test_cross_room_and_non_source_document_rejected(fix: ExtractionFixture) -> None:
    pdf_bytes = make_pdf(["Cross room"])
    asset_b = fix.asset_service.create(
        fix.owner_b,
        room_id=fix.room_b_id,
        kind="source_document",
        filename="room_b.pdf",
        mime_type="application/pdf",
        data=pdf_bytes,
        visibility="dm_only",
    )
    image_asset = fix.asset_service.create(
        fix.owner_a,
        room_id=fix.room_a_id,
        kind="image",
        filename="map.png",
        mime_type="image/png",
        data=b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15c4",
        visibility="room",
    )

    imp = fix.import_service.create_import(fix.owner_a, fix.room_a_id, "Reject Import")
    before_counts = _count_import_rows(fix.engine)

    # Cross-room asset rejected with not-found
    with pytest.raises(AdventureImportNotFoundError):
        fix.import_service.add_asset_source(
            fix.owner_a, fix.room_a_id, imp.id, asset_id=asset_b.id
        )
    assert _count_import_rows(fix.engine) == before_counts

    # Non-source-document asset (image) rejected with validation error
    with pytest.raises(AdventureImportValidationError):
        fix.import_service.add_asset_source(
            fix.owner_a, fix.room_a_id, imp.id, asset_id=image_asset.id
        )
    assert _count_import_rows(fix.engine) == before_counts


def test_media_mismatch_or_unsupported_rejected(fix: ExtractionFixture) -> None:
    imp = fix.import_service.create_import(fix.owner_a, fix.room_a_id, "Unsupported Media")
    before_counts = _count_import_rows(fix.engine)

    fake_asset = RoomAsset(
        id=uuid4(),
        room_id=fix.room_a_id,
        kind="source_document",
        original_filename="bad.xyz",
        mime_type="application/octet-stream",
        size_bytes=10,
        sha256="abc",
        visibility="dm_only",
        created_at=datetime.now(timezone.utc),
    )

    def _mock_open(*_args: object, **_kwargs: object):
        return fake_asset, BytesIO(b"0123456789")

    fix.asset_service.open_content = _mock_open  # type: ignore[method-assign]

    with pytest.raises(AdventureImportValidationError):
        fix.import_service.add_asset_source(
            fix.owner_a, fix.room_a_id, imp.id, asset_id=fake_asset.id
        )
    assert _count_import_rows(fix.engine) == before_counts


# --- 6. Text and Markdown Assets Reuse Normalization ---


@pytest.mark.parametrize(
    "source_kind,mime_type,filename,content,expected_norm",
    [
        (
            "txt",
            "text/plain",
            "notes.txt",
            b"  Line 1  \r\n\r\nLine 2  \r\n",
            "  Line 1\n\nLine 2\n",
        ),
        (
            "markdown",
            "text/markdown",
            "readme.md",
            b"# Heading\r\n\r\nParagraph text\r\n",
            "# Heading\n\nParagraph text\n",
        ),
    ],
)
def test_text_and_markdown_assets_reuse_normalization(
    fix: ExtractionFixture,
    source_kind: str,
    mime_type: str,
    filename: str,
    content: bytes,
    expected_norm: str,
) -> None:
    asset = fix.asset_service.create(
        fix.owner_a,
        room_id=fix.room_a_id,
        kind="source_document",
        filename=filename,
        mime_type=mime_type,
        data=content,
        visibility="dm_only",
    )

    imp = fix.import_service.create_import(fix.owner_a, fix.room_a_id, f"{source_kind.upper()} Asset")
    src = fix.import_service.add_asset_source(
        fix.owner_a, fix.room_a_id, imp.id, asset_id=asset.id
    )

    assert src.source_kind == source_kind
    assert src.asset_id == asset.id
    assert src.text_length == len(expected_norm)
    assert src.metadata_json["filename"] == filename
    assert src.metadata_json["media_type"] == mime_type

    # Chunk text is normalized
    chunk = fix.import_service.read_source_chunk(
        fix.owner_a, fix.room_a_id, imp.id, src.id
    )
    assert chunk.text == expected_norm

    # Re-adding is idempotent
    counts_before = _count_import_rows(fix.engine)
    src_again = fix.import_service.add_asset_source(
        fix.owner_a, fix.room_a_id, imp.id, asset_id=asset.id
    )
    counts_after = _count_import_rows(fix.engine)
    assert src_again.id == src.id
    assert counts_before == counts_after


# --- 7. Correction Regressions ---


class CloseTrackingBytesIO(BytesIO):
    def __init__(self, *args: object, **kwargs: object) -> None:
        super().__init__(*args, **kwargs)
        self.close_count = 0

    def close(self) -> None:
        self.close_count += 1
        super().close()


def test_asset_handle_closure_guaranteed_on_rejection(fix: ExtractionFixture) -> None:
    imp = fix.import_service.create_import(fix.owner_a, fix.room_a_id, "Handle Leak Test")

    # Case A: Reject due to kind != "source_document"
    bad_kind_asset = RoomAsset(
        id=uuid4(),
        room_id=fix.room_a_id,
        kind="image",
        original_filename="test.png",
        mime_type="image/png",
        size_bytes=10,
        sha256="abc",
        visibility="dm_only",
        created_at=datetime.now(timezone.utc),
    )
    tracking_handle_a = CloseTrackingBytesIO(b"dummy")

    def _mock_open_bad_kind(*_args: object, **_kwargs: object):
        return bad_kind_asset, tracking_handle_a

    fix.asset_service.open_content = _mock_open_bad_kind  # type: ignore[method-assign]

    with pytest.raises(AdventureImportValidationError):
        fix.import_service.add_asset_source(
            fix.owner_a, fix.room_a_id, imp.id, asset_id=bad_kind_asset.id
        )
    assert tracking_handle_a.closed is True
    assert tracking_handle_a.close_count >= 1

    # Case B: Reject due to unsupported mime_type
    bad_mime_asset = RoomAsset(
        id=uuid4(),
        room_id=fix.room_a_id,
        kind="source_document",
        original_filename="bad.xyz",
        mime_type="application/octet-stream",
        size_bytes=10,
        sha256="xyz",
        visibility="dm_only",
        created_at=datetime.now(timezone.utc),
    )
    tracking_handle_b = CloseTrackingBytesIO(b"dummy")

    def _mock_open_bad_mime(*_args: object, **_kwargs: object):
        return bad_mime_asset, tracking_handle_b

    fix.asset_service.open_content = _mock_open_bad_mime  # type: ignore[method-assign]

    with pytest.raises(AdventureImportValidationError):
        fix.import_service.add_asset_source(
            fix.owner_a, fix.room_a_id, imp.id, asset_id=bad_mime_asset.id
        )
    assert tracking_handle_b.closed is True
    assert tracking_handle_b.close_count >= 1


def test_docx_paragraph_index_preserves_gaps_from_empty_paragraphs() -> None:
    docx_bytes = make_docx(
        [
            ("Heading 1 Title", "Heading 1"),
            ("", "Normal"),
            ("Body paragraph 2", "Normal"),
            ("   \n   ", "Normal"),
            ("Heading 2 Subtitle", "Heading 2"),
        ]
    )

    result = extract_docx_text(docx_bytes)
    assert len(result.sections) == 3

    # Preserves original doc.paragraphs index (0, 2, 4), not 0, 1, 2
    s0, s1, s2 = result.sections
    assert s0["paragraph_index"] == 0
    assert s0["heading_index"] == 0
    assert result.normalized_text[int(s0["start_offset"]) : int(s0["end_offset"])] == "Heading 1 Title"

    assert s1["paragraph_index"] == 2
    assert s1["heading_index"] is None
    assert result.normalized_text[int(s1["start_offset"]) : int(s1["end_offset"])] == "Body paragraph 2"

    assert s2["paragraph_index"] == 4
    assert s2["heading_index"] == 1
    assert result.normalized_text[int(s2["start_offset"]) : int(s2["end_offset"])] == "Heading 2 Subtitle"


def test_extractors_preserve_normalized_leading_spaces() -> None:
    # PDF: preserve leading indentation
    pdf_bytes = make_pdf(["  Indented line 1\n    Indented line 2"])
    res_pdf = extract_pdf_text(pdf_bytes)
    assert res_pdf.normalized_text.startswith("  Indented line 1")
    assert "    Indented line 2" in res_pdf.normalized_text
    assert res_pdf.sections[0]["start_offset"] == 0

    # DOCX: preserve leading indentation
    docx_bytes = make_docx([("  Indented paragraph", "Normal")])
    res_docx = extract_docx_text(docx_bytes)
    assert res_docx.normalized_text.startswith("  Indented paragraph")
    assert res_docx.sections[0]["start_offset"] == 0


def test_add_asset_source_authority_rejections(fix: ExtractionFixture) -> None:
    imp = fix.import_service.create_import(fix.owner_a, fix.room_a_id, "Auth Test")
    pdf_bytes = make_pdf(["Content"])
    asset = fix.asset_service.create(
        fix.owner_a,
        room_id=fix.room_a_id,
        kind="source_document",
        filename="auth.pdf",
        mime_type="application/pdf",
        data=pdf_bytes,
        visibility="dm_only",
    )

    before_counts = _count_import_rows(fix.engine)

    # Member in the room rejected with typed forbidden error
    with pytest.raises(AdventureImportForbiddenError):
        fix.import_service.add_asset_source(
            fix.member_a, fix.room_a_id, imp.id, asset_id=asset.id
        )
    assert _count_import_rows(fix.engine) == before_counts

    # Owner of another room rejected with typed not-found error
    with pytest.raises(AdventureImportNotFoundError):
        fix.import_service.add_asset_source(
            fix.owner_b, fix.room_a_id, imp.id, asset_id=asset.id
        )
    assert _count_import_rows(fix.engine) == before_counts
